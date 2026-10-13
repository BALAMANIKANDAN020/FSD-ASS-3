import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Page setup
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Document Header / Title
title_p = doc.add_paragraph()
title_run = title_p.add_run("Coding assignment 3 solu")
title_run.font.size = Pt(20)
title_run.font.bold = True
title_run.font.name = "Calibri"
title_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle_p = doc.add_paragraph()
sub_run = subtitle_p.add_run("Containerize And Deploy An App")
sub_run.font.size = Pt(13)
sub_run.font.italic = True
sub_run.font.name = "Calibri"
sub_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()

# Learning outcome
p_lo = doc.add_paragraph()
r_lo_label = p_lo.add_run("Learning Outcome Addressed: ")
r_lo_label.font.bold = True
r_lo_label.font.size = Pt(11)
r_lo_val = p_lo.add_run("4. Deploy a dockerized application to the cloud.")
r_lo_val.font.size = Pt(11)

def add_code_block(document, title, code_str):
    h = document.add_heading(title, level=1)
    p = document.add_paragraph()
    p_run = p.add_run(code_str)
    p_run.font.name = "Consolas"
    p_run.font.size = Pt(9)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)

base_dir = os.path.dirname(os.path.abspath(__file__))

# 1. Dockerfile
with open(os.path.join(base_dir, "Dockerfile"), "r", encoding="utf-8") as f:
    dockerfile_code = f.read()
add_code_block(doc, "Dockerfile", dockerfile_code)

# 2. .dockerignore
with open(os.path.join(base_dir, ".dockerignore"), "r", encoding="utf-8") as f:
    dockerignore_code = f.read()
add_code_block(doc, ".dockerignore", dockerignore_code)

# 3. package.json
with open(os.path.join(base_dir, "package.json"), "r", encoding="utf-8") as f:
    package_code = f.read()
add_code_block(doc, "package.json", package_code)

# 4. App.js
with open(os.path.join(base_dir, "src", "App.js"), "r", encoding="utf-8") as f:
    app_js_code = f.read()
add_code_block(doc, "App.js", app_js_code)

# 5. Commands Section
commands = """=== OPTION 1: HEROKU CONTAINER REGISTRY ===
# 1. Build Docker image locally
docker build -t my-react-app .

# 2. Verify local Docker images
docker images

# 3. Run container locally on port 3000
docker run -p 3000:3000 my-react-app
# Open in browser: http://localhost:3000

# 4. Heroku CLI Login & Remote Link
heroku login
heroku git:remote -a <your-app-name>

# 5. Log in to Heroku Container Registry
heroku container:login

# 6. Push Docker image to Heroku Container Registry
heroku container:push web -a <your-app-name>

# 7. Release container to live web dyno
heroku container:release web -a <your-app-name>

# 8. Open live cloud application
heroku open -a <your-app-name>

=== OPTION 2: RENDER CLOUD DOCKER DEPLOYMENT (FREE TIER) ===
# 1. Push project to GitHub
git remote add origin https://github.com/balamanikandan020/<repo-name>.git
git branch -M main
git push -u origin main

# 2. On Render Dashboard (https://dashboard.render.com):
#    - Click "New +" -> "Web Service"
#    - Connect your GitHub repository
#    - Runtime: Docker (Automatically builds Dockerfile)
#    - Plan: Free
#    - Click "Deploy Web Service"
#    - Live URL: https://<app-name>.onrender.com"""
add_code_block(doc, "Step-by-Step Terminal Commands", commands)

# 6. Output & Screenshots Checklist
doc.add_heading("Output & Screenshots Verification", level=1)

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "Step #"
hdr_cells[1].text = "Screenshot Description"
hdr_cells[2].text = "Status / Purpose"

for cell in hdr_cells:
    for par in cell.paragraphs:
        for run in par.runs:
            run.font.bold = True
            run.font.size = Pt(10)

items = [
    ("1", "GitHub Repository created with Dockerfile", "GitHub repo push successful"),
    ("2", "Docker image build successful (terminal/cloud)", "docker build -t my-react-app ."),
    ("3", "Local browser test (http://localhost:3000)", "Working React application locally"),
    ("4", "Cloud Dashboard (Heroku / Render)", "Cloud Web Service created from Dockerfile"),
    ("5", "Docker container build & deploy logs", "Successful container build in cloud"),
    ("6", "Final Live Deployed Cloud Application ⭐", "Live React App on Cloud (Heroku / Render)"),
]

for row_data in items:
    row_cells = table.add_row().cells
    row_cells[0].text = row_data[0]
    row_cells[1].text = row_data[1]
    row_cells[2].text = row_data[2]
    for cell in row_cells:
        for par in cell.paragraphs:
            for run in par.runs:
                run.font.size = Pt(9.5)

doc.add_paragraph()

p_note = doc.add_paragraph()
r_note = p_note.add_run("Note: Paste your final screenshots below each corresponding section as required by the assignment evaluation guidelines.")
r_note.font.italic = True
r_note.font.size = Pt(10)

doc.add_paragraph()
doc.add_paragraph()

# Student details
p_student = doc.add_paragraph()
p_student_run = p_student.add_run("BALAMANIKANDAN R (1P25MC006)")
p_student_run.font.size = Pt(13)
p_student_run.font.bold = True
p_student_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
p_student.alignment = WD_ALIGN_PARAGRAPH.RIGHT

target_output = r"d:\MCA-3\FULL STACK DEVELOPER\CODING ACTIVITY\UNIT - 3\CODING ASSIGNMENT-3.docx"
try:
    doc.save(target_output)
    print(f"Saved {target_output} successfully!")
except PermissionError:
    alt_output = r"d:\MCA-3\FULL STACK DEVELOPER\CODING ACTIVITY\UNIT - 3\CODING ASSIGNMENT-3_COMPLETED.docx"
    doc.save(alt_output)
    print(f"Target file was open in Word. Saved {alt_output} successfully!")
